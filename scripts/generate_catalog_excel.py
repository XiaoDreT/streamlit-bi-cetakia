import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

def create_bi_insight_catalog():
    wb = openpyxl.Workbook()
    # remove default sheet
    wb.remove(wb.active)

    # Styling definitions
    font_family = "Calibri"
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    title_font = Font(name=font_family, size=14, bold=True, color="1F4E78")
    data_font = Font(name=font_family, size=10, color="000000")
    bold_data_font = Font(name=font_family, size=10, bold=True, color="000000")
    
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    alt_fill = PatternFill(start_color="F2F6FA", end_color="F2F6FA", fill_type="solid")
    
    thin_border_side = Side(border_style="thin", color="D9D9D9")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    # -------------------------------------------------------------
    # SHEET 1: Dashboard Overview
    # -------------------------------------------------------------
    ws1 = wb.create_sheet(title="Dashboard Overview")
    ws1.views.sheetView[0].showGridLines = True
    
    sheet1_headers = ["Dashboard Name", "User", "Purpose", "Priority"]
    sheet1_data = [
        ["Executive Business Dashboard", "Owner, Executive Management", "High-level business health monitoring, revenue growth drivers, repeat vs new source decomposition, branch benchmarking, top customer concentration risk exposure.", "High (P0)"],
        ["Customer Intelligence Dashboard", "Customer Service Manager, Marketing Lead", "Customer health distribution (Healthy, Active, At Risk, Dormant), RFM segmentation grid, high-value churn risk worklist, 12-month cohort retention heatmap.", "High (P0)"],
        ["Customer 360 Dashboard", "Customer Service Staff, Sales Staff, Credit Analyst", "Single-view 360 profile per account, historical purchase velocity, preferred product basket, cross-sell recommendations, financial payment compliance class.", "High (P0)"],
        ["Sales Intelligence Dashboard", "Sales Manager, Sales Staff, Commercial Ops", "Multi-stage quotation conversion funnel, expired quotation recovery pipeline leaderboard, sales representative performance matrix, urgent expiry worklist.", "High (P0)"],
        ["Product Intelligence Dashboard", "Marketing Lead, Product Category Manager", "Category revenue/volume mix, high-revenue low-penetration quadrant matrix, basket co-purchase affinity pairs (Support/Confidence/Lift), bundle engine.", "Medium (P1)"],
        ["Market Intelligence Dashboard", "Owner, Marketing Director, Business Dev", "Customer segment revenue contribution (Instansi, End User, Agen, Sekolah, UMKM, Industri, Employee, Divisi), regional branch catchment benchmarking across 6 branches (Cipta Graha, Cipta Digital, Cipta Galuh, Cipta Cianjur, Cipta Purwakarta, Cipta Online), acquisition velocity, cluster density.", "Medium (P1)"]
    ]

    ws1.append(sheet1_headers)
    for row_idx, row_data in enumerate(sheet1_data, start=2):
        ws1.append(row_data)

    # -------------------------------------------------------------
    # SHEET 2: Insight Catalog
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Insight Catalog")
    ws2.views.sheetView[0].showGridLines = True

    sheet2_headers = [
        "Insight ID", "Insight Name", "Dashboard Module", "User Role",
        "Business Question", "Dataset Source", "Required Columns", "Formula",
        "Visualization Type", "Priority", "Business Meaning", "Recommended Action"
    ]

    sheet2_data = [
        # Executive Module
        [
            "INS-EXEC-01", "Net Sales Performance & Growth", "Executive Business Dashboard", "Owner, Management",
            "What is total net revenue and how does it compare to prior period?",
            "02_sales_transaction_extract.csv", "net_sales, invoice_date, division_name",
            "SUM(net_sales) [Current Period] vs SUM(net_sales) [Prior Period]",
            "KPI Metric Card + Sparkline", "High (P0)",
            "Evaluates macro revenue trajectory and delta against baseline period across 6 regional branches.",
            "Owner: Investigate volume vs value drivers before adjusting targets."
        ],
        [
            "INS-EXEC-02", "Revenue Source Breakdown", "Executive Business Dashboard", "Owner, Management",
            "Is revenue coming from new customer acquisition or repeat spend?",
            "02_sales_transaction_extract.csv, 04_customer_intelligence_analysis.csv", "net_sales, customer_id, first_purchase_date",
            "SUM(net_sales) WHERE invoice_date == first_purchase_date vs SUM(net_sales) WHERE invoice_date > first_purchase_date",
            "Stacked Bar Chart", "High (P0)",
            "Measures company reliance on repeat customer retention vs acquisition.",
            "Marketing: Launch acquisition campaigns if new revenue share < 25%."
        ],
        [
            "INS-EXEC-03", "Branch Revenue Benchmarking", "Executive Business Dashboard", "Owner, Branch Lead",
            "How is revenue and active customer count distributed across the 6 regional branches (Cipta Graha, Cipta Digital, Cipta Galuh, Cipta Cianjur, Cipta Purwakarta, Cipta Online)?",
            "02_sales_transaction_extract.csv", "net_sales, invoice_id, division_name, customer_id",
            "SUM(net_sales), COUNT(DISTINCT invoice_id), COUNT(DISTINCT customer_id) GROUP BY division_name",
            "Grouped Column Chart", "Medium (P1)",
            "Highlights regional capacity usage and market share differences across 6 Cipta Grafika branches.",
            "Management: Reallocate regional sales support to lagging branches based on capacity."
        ],
        [
            "INS-EXEC-04", "Top Customer Concentration & Health", "Executive Business Dashboard", "Owner, CS Manager",
            "Are top revenue-generating customers healthy or at risk of churning?",
            "02_sales_transaction_extract.csv, 04_customer_intelligence_analysis.csv", "customer_id, customer_name, net_sales, customer_health",
            "TOP 10 customer_id BY SUM(net_sales) JOIN customer_health",
            "Leaderboard Table", "High (P0)",
            "Uncovers revenue risk exposure if key accounts slip to At Risk/Dormant.",
            "Owner: Executive check-in calls for top accounts flagged At Risk."
        ],

        # Customer Intelligence Module
        [
            "INS-CUST-01", "Customer Health Portfolio Distribution", "Customer Intelligence Dashboard", "CS Manager, Marketing",
            "What proportion of our customer base is Healthy, Active, At Risk, or Dormant?",
            "04_customer_intelligence_analysis.csv", "customer_id, customer_health, days_since_last_purchase",
            "COUNT(customer_id) GROUP BY customer_health",
            "Donut Chart", "High (P0)",
            "Quantifies migration of customer accounts into churn risk categories.",
            "CS Lead: Target At Risk accounts with lifetime spend > Rp 10M."
        ],
        [
            "INS-CUST-02", "High-Value At-Risk Customer Worklist", "Customer Intelligence Dashboard", "CS Staff, Account Manager",
            "Which specific high-value accounts have exceeded their expected purchase cycle?",
            "04_customer_intelligence_analysis.csv", "customer_id, customer_name, lifetime_sales, days_since_last_purchase, customer_health",
            "FILTER customer_health IN ('At Risk', 'Dormant') SORT BY lifetime_sales DESC",
            "Worklist Table", "High (P0)",
            "Provides immediate outbound retention contact target list.",
            "CS Staff: Execute outbound wellness call within 48 hours."
        ],
        [
            "INS-CUST-03", "Multi-Month Cohort Retention", "Customer Intelligence Dashboard", "Marketing Lead, CS Manager",
            "How well do new customer acquisition cohorts retain over 12 months?",
            "04_customer_intelligence_analysis.csv, 02_sales_transaction_extract.csv", "customer_id, first_purchase_date, invoice_date",
            "(COUNT(Active Customers in Month N) / Initial Cohort Size in Month 0) * 100",
            "Cohort Heatmap", "Medium (P1)",
            "Evaluates long-term customer lifecycle retention decay.",
            "Marketing: Redesign onboarding experience if M+1 retention < 35%."
        ],
        [
            "INS-CUST-04", "RFM Segment Composition", "Customer Intelligence Dashboard", "Marketing Lead, Sales",
            "How are customers distributed across RFM loyalty tiers (Champions to Hibernating)?",
            "04_customer_intelligence_analysis.csv", "customer_id, R_score, F_score, M_score, rfm_segment",
            "COUNT(customer_id), SUM(lifetime_sales) GROUP BY rfm_segment",
            "Treemap / Bubble Grid", "Medium (P1)",
            "Maps customer base quality and loyalty concentration.",
            "Marketing: Tailor campaign offers based on RFM segment behavior."
        ],

        # Customer 360 Module
        [
            "INS-C360-01", "Single Customer Lifetime Profile", "Customer 360 Dashboard", "CS Staff, Sales Rep",
            "What is the total lifetime spend, order frequency, and health status of account X?",
            "01_customer_master_raw.csv, 04_customer_intelligence_analysis.csv", "customer_id, lifetime_sales, total_orders, average_order_value, customer_health",
            "SELECT lifetime_sales, total_orders, average_order_value, customer_health WHERE customer_id == X",
            "Profile Header Card", "High (P0)",
            "Single view of customer relationship value and health status.",
            "CS/Sales: Adjust negotiation terms based on lifetime tier."
        ],
        [
            "INS-C360-02", "Purchase Velocity & Reorder Gap", "Customer 360 Dashboard", "CS Staff, Sales Rep",
            "Is customer X due for a reorder based on their historical buying interval?",
            "02_sales_transaction_extract.csv, 04_customer_intelligence_analysis.csv", "customer_id, invoice_date, days_since_last_purchase",
            "AVG(DATEDIFF(invoice_date_t, invoice_date_{t-1})) vs days_since_last_purchase",
            "Velocity Gauge", "High (P0)",
            "Predicts natural reorder timing for proactive outreach.",
            "Sales: Contact customer for reorder if recency > historical avg gap."
        ],
        [
            "INS-C360-03", "Preferred Product Basket & Cross-Sell", "Customer 360 Dashboard", "Sales Rep, CS Staff",
            "What are customer X's top purchased products and recommended next purchases?",
            "03_sales_item_extract.csv", "customer_id, product_id, product_name, sales_amount",
            "SUM(sales_amount) GROUP BY product_id ORDER BY SUM(sales_amount) DESC",
            "Horizontal Bar Chart", "Medium (P1)",
            "Highlights wallet-share expansion potential per account.",
            "Sales: Pitch complementary co-purchase bundle during quote draft."
        ],
        [
            "INS-C360-04", "Financial Compliance & Payment Class", "Customer 360 Dashboard", "Credit Analyst, Finance",
            "Does customer X have outstanding overdue balances or payment risk ratings?",
            "02_sales_transaction_extract.csv, 04_customer_intelligence_analysis.csv", "customer_id, grand_total, invoice_status",
            "SUM(grand_total) WHERE invoice_status == 'Unpaid' AND due_date < as_of",
            "Financial Risk Widget", "High (P0)",
            "Protects liquidity by preventing credit extensions to bad payers.",
            "Credit Analyst: Enforce Cash-on-Delivery if overdue > 30 days."
        ],

        # Sales Intelligence Module
        [
            "INS-SALE-01", "Multi-Stage Quotation Conversion Funnel", "Sales Intelligence Dashboard", "Sales Manager, Commercial Lead",
            "What percentage of issued quotations convert into confirmed sales orders and invoices?",
            "05_quotation_funnel_master.csv", "quotation_id, conversion_status, converted_flag, net_quotation_value",
            "(COUNT(quotation_id WHERE converted_flag == 1) / TOTAL quotation_id) * 100",
            "Funnel Chart", "High (P0)",
            "Measures pipeline throughput and identifies stage drop-offs.",
            "Sales Manager: Review sales pitch quality if Stage 1->2 drop > 50%."
        ],
        [
            "INS-SALE-02", "Expired & Lost Opportunity Pipeline", "Sales Intelligence Dashboard", "Sales Manager, Sales Rep",
            "What is the total value of expired quotations, and who are the assigned sales reps?",
            "05_quotation_funnel_master.csv", "quotation_id, unconverted_quotation_value, quotation_status, sales_name",
            "SUM(unconverted_quotation_value) WHERE quotation_status == 'Expired' GROUP BY sales_name",
            "Bar Chart + Worklist", "High (P0)",
            "Quantifies uncaptured pipeline revenue eligible for recovery.",
            "Sales Manager: Enforce 7-day post-expiry call follow-up campaign."
        ],
        [
            "INS-SALE-03", "Sales Representative Performance Matrix", "Sales Intelligence Dashboard", "Sales Manager",
            "How do sales reps compare in terms of revenue billed, win rate, and average deal size?",
            "05_quotation_funnel_master.csv, 02_sales_transaction_extract.csv", "sales_code, sales_name, net_sales, converted_flag, net_quotation_value",
            "SUM(net_sales), WIN_RATE = (Converted / Total Quotes), AOV = (Sales / Orders)",
            "Scatter Plot Matrix", "Medium (P1)",
            "Benchmarks sales efficiency and individual closing ability.",
            "Sales Manager: Pair low-conversion reps with top closers for mentorship."
        ],
        [
            "INS-SALE-04", "Urgent Quotation Expiry Alert", "Sales Intelligence Dashboard", "Sales Rep",
            "Which open quotations are set to expire within 72 hours?",
            "05_quotation_funnel_master.csv", "quotation_id, quotation_code, days_from_expiry, net_quotation_value, sales_name",
            "FILTER quotation_status == 'Unconverted' AND days_from_expiry <= 3",
            "Urgent Alert Worklist", "High (P0)",
            "Prevents quote expiration through timely sales intervention.",
            "Sales Rep: Immediate phone follow-up with client before expiration."
        ],

        # Product Intelligence Module
        [
            "INS-PROD-01", "Product Category Revenue Contribution", "Product Intelligence Dashboard", "Marketing Lead, Category Manager",
            "Which product categories generate the highest revenue and order volume?",
            "03_sales_item_extract.csv", "product_category, sales_amount, qty",
            "SUM(sales_amount), SUM(qty), SHARE = (Category Sales / Total Sales) * 100",
            "Donut Chart + Table", "Medium (P1)",
            "Identifies core category anchors vs niche product lines.",
            "Category Manager: Adjust raw material procurement based on volume."
        ],
        [
            "INS-PROD-02", "High Revenue / Low Penetration Opportunity", "Product Intelligence Dashboard", "Marketing Lead, Sales",
            "Which high-revenue products are purchased by only a small fraction of active customers?",
            "03_sales_item_extract.csv, 02_sales_transaction_extract.csv", "product_id, product_name, sales_amount, customer_id",
            "Penetration = (COUNT(DISTINCT customer_id buying Product) / Total Active Customers) * 100",
            "4-Quadrant Scatter Plot", "High (P0)",
            "Highlights star products with massive cross-sell potential.",
            "Marketing: Launch product expansion campaign to non-buying base."
        ],
        [
            "INS-PROD-03", "Co-Purchase Basket Affinity (Pairs)", "Product Intelligence Dashboard", "Marketing Lead, Sales Rep",
            "Which product pairs are statistically co-purchased together in single orders?",
            "03_sales_item_extract.csv", "invoice_id, product_id, product_category",
            "Support(A,B) = N_AB/N, Confidence(A->B) = N_AB/N_A, Lift = Conf / Support(B)",
            "Co-purchase Matrix Table", "Medium (P1)",
            "Provides statistical evidence of product co-buying patterns.",
            "Sales Rep: Offer 5% bundle discount when customer buys Product A."
        ],
        [
            "INS-PROD-04", "Product Bundle Recommendation Engine", "Product Intelligence Dashboard", "Marketing Lead",
            "What commercial product packages should be created for specific customer segments (Instansi, End User, Agen, Sekolah, UMKM, Industri, Employee, Divisi)?",
            "03_sales_item_extract.csv, 01_customer_master_raw.csv", "product_category, customer_category, invoice_id",
            "Identify top Lift pairs grouped by customer_category",
            "Recommendation Card", "Medium (P1)",
            "Packages affinity pairs into ready-to-sell commercial bundles.",
            "Marketing: Publish UMKM Starter Pack and Industri Packaging bundle in Cetakia order entry."
        ],

        # Market Intelligence Module
        [
            "INS-MKT-01", "Customer Segment Market Contribution", "Market Intelligence Dashboard", "Owner, Marketing Director",
            "How is revenue and customer population divided across Instansi, End User, Agen, Sekolah, UMKM, Industri, Employee, and Divisi?",
            "01_customer_master_raw.csv, 02_sales_transaction_extract.csv", "customer_category, net_sales, customer_id",
            "SUM(net_sales), COUNT(DISTINCT customer_id) GROUP BY customer_category",
            "Treemap Chart", "Medium (P1)",
            "Assesses market concentration and segment diversification across 8 customer categories.",
            "Executive: Assign Key Account Manager team if Industri > 40% revenue."
        ],
        [
            "INS-MKT-02", "Regional Branch Market Catchment", "Market Intelligence Dashboard", "Owner, Branch Manager",
            "How effectively is each of the 6 regional branches (Cipta Graha, Cipta Digital, Cipta Galuh, Cipta Cianjur, Cipta Purwakarta, Cipta Online) penetrating its regional prospect market?",
            "01_customer_master_raw.csv, 02_sales_transaction_extract.csv", "division_name, prospect, customer_id, net_sales",
            "COUNT(customer_id WHERE prospect == False) vs COUNT(customer_id WHERE prospect == True)",
            "Grouped Bar Chart", "Medium (P1)",
            "Benchmarks branch sales coverage against local potential across 6 regional branches.",
            "Marketing: Deploy local lead generation in lagging branch areas (Cianjur, Online, Purwakarta)."
        ],
        [
            "INS-MKT-03", "Customer Acquisition & First Purchase Speed", "Market Intelligence Dashboard", "Marketing Lead",
            "How fast do newly registered customers convert into their first paid order?",
            "01_customer_master_raw.csv, 04_customer_intelligence_analysis.csv", "customer_created_at, first_purchase_date",
            "AVG(DATEDIFF(first_purchase_date, customer_created_at))",
            "Line Chart", "Medium (P1)",
            "Measures onboarding funnel efficiency and speed to value.",
            "Marketing: Issue first-order welcome discount if speed > 14 days."
        ],
        [
            "INS-MKT-04", "Industry Cluster Revenue Density", "Market Intelligence Dashboard", "Marketing Lead",
            "Which industry sub-clusters (e.g. F&B, Pharma, Retail) show the highest spend density?",
            "01_customer_master_raw.csv, 04_customer_intelligence_analysis.csv", "cluster, lifetime_sales, customer_id",
            "SUM(lifetime_sales), COUNT(customer_id) GROUP BY cluster",
            "Bubble Chart", "Medium (P1)",
            "Uncovers niche industry specialization opportunities.",
            "Product: Create industry-specific print spec sample kits."
        ]
    ]

    ws2.append(sheet2_headers)
    for row in sheet2_data:
        ws2.append(row)

    # -------------------------------------------------------------
    # SHEET 3: Metric Dictionary
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Metric Dictionary")
    ws3.views.sheetView[0].showGridLines = True

    sheet3_headers = ["Metric Name", "Definition", "Formula", "Data Source", "Business Owner"]
    sheet3_data = [
        [
            "Net Sales Amount",
            "Total net monetary revenue generated from confirmed invoice transactions after discounts.",
            "SUM(net_sales)",
            "02_sales_transaction_extract.csv",
            "Owner / Finance Lead"
        ],
        [
            "Order Count",
            "Total number of distinct invoice transactions generated during the specified period.",
            "COUNT(DISTINCT invoice_id)",
            "02_sales_transaction_extract.csv",
            "Sales Manager"
        ],
        [
            "Active Customers",
            "Count of unique customer accounts that completed at least 1 purchase within the period.",
            "COUNT(DISTINCT customer_id)",
            "02_sales_transaction_extract.csv",
            "Customer Service Manager"
        ],
        [
            "Average Order Value (AOV)",
            "Mean net monetary revenue earned per invoice transaction.",
            "SUM(net_sales) / COUNT(DISTINCT invoice_id)",
            "02_sales_transaction_extract.csv",
            "Sales Manager"
        ],
        [
            "Customer Recency",
            "Elapsed time in calendar days between cutoff date (as_of) and customer's latest purchase.",
            "as_of_date - MAX(invoice_date)",
            "04_customer_intelligence_analysis.csv",
            "Customer Service Manager"
        ],
        [
            "Reorder Interval Cycle",
            "Historical average gap in calendar days between consecutive purchases for a customer.",
            "AVG(DATEDIFF(order_t, order_{t-1}))",
            "04_customer_intelligence_analysis.csv",
            "Customer Service Manager"
        ],
        [
            "Customer Health Status",
            "Categorical state (Healthy, Active, At Risk, Dormant, New) derived from recency vs cycle.",
            "Rule Engine: At Risk if Recency > MAX(30d, 1.5x Cycle); Dormant if Recency > MAX(90d, 3x Cycle)",
            "04_customer_intelligence_analysis.csv",
            "Customer Service Manager"
        ],
        [
            "Quotation Conversion Rate",
            "Percentage of issued quotation cohort that successfully converted into sales orders.",
            "(COUNT(Converted Quotations) / COUNT(Total Quotation Cohort)) * 100",
            "05_quotation_funnel_master.csv",
            "Sales Manager"
        ],
        [
            "Expired Quotation Value",
            "Total uncaptured revenue trapped in quotations that passed expiry date without conversion.",
            "SUM(unconverted_quotation_value WHERE quotation_status == 'Expired')",
            "05_quotation_funnel_master.csv",
            "Sales Manager"
        ],
        [
            "Product Penetration Rate",
            "Proportion of active customers who bought a specific product SKU within the period.",
            "(COUNT(Unique Product Buyers) / COUNT(Total Active Customers in Scope)) * 100",
            "03_sales_item_extract.csv",
            "Marketing Lead"
        ],
        [
            "Co-Purchase Lift",
            "Statistical measure of how much more often Product A and B are bought together than expected.",
            "Confidence(A->B) / Support(B) = (N_AB * N_Total) / (N_A * N_B)",
            "03_sales_item_extract.csv",
            "Marketing Lead"
        ],
        [
            "Cohort Retention Rate M+n",
            "Percentage of initial acquisition cohort customers who make a repeat purchase in Month n.",
            "(COUNT(Active Buyers in Month n) / Initial Cohort Size in Month 0) * 100",
            "04_customer_intelligence_analysis.csv",
            "Marketing Lead"
        ],
        [
            "Cash Collection Rate",
            "Percentage of billed invoice monetary value settled and paid as of the cutoff date.",
            "(SUM(Allocated Payment Amount) / SUM(Grand Total Billed)) * 100",
            "02_sales_transaction_extract.csv",
            "Finance Lead"
        ],
        [
            "Outstanding Receivables",
            "Total unpaid monetary invoice balance remaining as of the cutoff date.",
            "SUM(grand_total - allocated_payments)",
            "02_sales_transaction_extract.csv",
            "Finance Lead"
        ],
        [
            "Lifetime Observed Sales",
            "Cumulative net sales revenue billed to a customer across the available 12-month history.",
            "SUM(net_sales WHERE customer_id == X)",
            "04_customer_intelligence_analysis.csv",
            "Customer Service Manager"
        ]
    ]

    ws3.append(sheet3_headers)
    for row in sheet3_data:
        ws3.append(row)

    # -------------------------------------------------------------
    # SHEET 4: API Requirement
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="API Requirement")
    ws4.views.sheetView[0].showGridLines = True

    sheet4_headers = ["Insight ID", "API Endpoint", "Request Parameter", "Response Field", "Refresh Frequency"]
    sheet4_data = [
        ["INS-EXEC-01", "/api/bi/dashboard-summary", "date_from, date_to, branch_ids, timezone", "sales_amount.value, previous, growth_pct, comparison_status", "Hourly"],
        ["INS-EXEC-02", "/api/bi/dashboard-summary", "date_from, date_to, branch_ids", "revenue_sources.first_purchase, repeat_purchase", "Hourly"],
        ["INS-EXEC-03", "/api/bi/sales/performance", "date_from, date_to, group_by=branch", "branch_id, division_name, net_sales, order_count, active_customers", "Hourly"],
        ["INS-EXEC-04", "/api/bi/dashboard-summary", "date_from, date_to, limit=10", "top_customers.customer_id, customer_name, net_sales, health_status", "Hourly"],
        ["INS-CUST-01", "/api/bi/customers/health", "as_of, branch_ids, segments", "health_summary.healthy, active, at_risk, dormant, new", "Daily"],
        ["INS-CUST-02", "/api/bi/customers/health", "as_of, status=at_risk, dormant, sort=lifetime_sales", "items[].customer_id, customer_name, lifetime_sales, recency_days", "Daily"],
        ["INS-CUST-03", "/api/bi/customers/cohorts", "as_of, cohort_year=2025", "cohorts[].cohort_month, initial_size, retention_matrix[]", "Weekly"],
        ["INS-CUST-04", "/api/bi/customers/segments", "as_of, rfm=true", "rfm_distribution[].segment_name, count, total_lifetime_sales", "Daily"],
        ["INS-C360-01", "/api/bi/customer/{id}", "customer_id, as_of", "customer_profile.name, category, lifetime_sales, health_status, rfm_score", "Real-time"],
        ["INS-C360-02", "/api/bi/customer/{id}", "customer_id, as_of", "order_velocity.recency_days, avg_reorder_gap_days, next_expected_order", "Real-time"],
        ["INS-C360-03", "/api/bi/customer/{id}", "customer_id, top=5", "top_products[].product_name, spend, cross_sell_recommendations[]", "Real-time"],
        ["INS-C360-04", "/api/bi/customer/{id}", "customer_id, as_of", "financial_status.billed_total, outstanding_balance, payment_class", "Real-time"],
        ["INS-SALE-01", "/api/bi/sales/funnel", "date_from, date_to, branch_ids", "stages[].stage_name, quote_count, value, conversion_rate", "Hourly"],
        ["INS-SALE-02", "/api/bi/sales/quotations", "status=expired, date_from, date_to", "expired_summary.total_value, items[].sales_name, customer_name, value", "Hourly"],
        ["INS-SALE-03", "/api/bi/sales/performance", "date_from, date_to, group_by=salesperson", "reps[].sales_code, sales_name, net_sales, win_rate, average_deal_size", "Hourly"],
        ["INS-SALE-04", "/api/bi/sales/quotations", "status=unconverted, max_days_expiry=3", "urgent_items[].quotation_code, customer_name, value, days_from_expiry", "Real-time"],
        ["INS-PROD-01", "/api/bi/products/top", "date_from, date_to, group_by=category", "categories[].category_name, sales_amount, qty_sold, share_pct", "Daily"],
        ["INS-PROD-02", "/api/bi/products/opportunities", "date_from, date_to, penetration_max=20", "opportunities[].product_name, sales_amount, buyer_count, penetration_rate", "Daily"],
        ["INS-PROD-03", "/api/bi/products/pairs", "date_from, date_to, min_count=10", "pairs[].product_a, product_b, co_purchase_count, support, confidence, lift", "Weekly"],
        ["INS-PROD-04", "/api/bi/products/pairs", "date_from, date_to, bundle_mode=true", "bundles[].bundle_name, products[], lift, target_segment", "Weekly"],
        ["INS-MKT-01", "/api/bi/customers/segments", "date_from, date_to", "segments[].category_name, net_sales, customer_count, avg_ltv", "Daily"],
        ["INS-MKT-02", "/api/bi/sales/performance", "date_from, date_to, include_prospects=true", "branches[].division_name, active_buyers, unconverted_prospects", "Daily"],
        ["INS-MKT-03", "/api/bi/customers/segments", "acquisition_trends=true", "monthly_acquisition[].month, registrations, first_purchases, avg_speed_days", "Weekly"],
        ["INS-MKT-04", "/api/bi/customers/segments", "group_by=cluster", "clusters[].cluster_name, total_sales, customer_count, avg_order_val", "Weekly"]
    ]

    ws4.append(sheet4_headers)
    for row in sheet4_data:
        ws4.append(row)

    # -------------------------------------------------------------
    # FORMATTING ALL SHEETS
    # -------------------------------------------------------------
    for ws in wb.worksheets:
        # Format Header Row
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = thin_border

        # Format Data Rows
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            fill = alt_fill if row_idx % 2 == 0 else PatternFill(fill_type=None)
            for cell in row:
                cell.font = data_font
                cell.border = thin_border
                if fill.fill_type:
                    cell.fill = fill
                
                # Alignments based on sheet & column
                col_name = ws.cell(row=1, column=cell.column).value
                if col_name in ["Insight ID", "Priority", "Refresh Frequency", "User Role", "Dashboard Name"]:
                    cell.alignment = align_center
                elif col_name in ["Formula", "Required Columns", "Response Field", "Request Parameter", "API Endpoint"]:
                    cell.alignment = align_left
                else:
                    cell.alignment = align_left

        # Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                # If multi-line or very long, cap length for width calc
                lines = val_str.split('\n')
                for line in lines:
                    if len(line) > max_len:
                        max_len = len(line)
            # Apply padding & column constraints
            width = min(max(max_len + 4, 12), 45)
            ws.column_dimensions[col_letter].width = width

    output_dir = "/home/zen/Documents/Cetakia/cetakia-bi/docs"
    os.makedirs(output_dir, exist_ok=True)
    file_path = os.path.join(output_dir, "BI_Insight_Catalog_V1.xlsx")
    wb.save(file_path)
    print(f"Successfully generated Excel workbook at: {file_path}")

if __name__ == "__main__":
    create_bi_insight_catalog()
